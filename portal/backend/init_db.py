import sqlite3
import openpyxl
import os
import json
from datetime import datetime

EXCEL_PATH = r"d:\AI project\Production & Store Plan\SKU_RM_Planning_System.xlsx"
DB_PATH = os.path.join(os.path.dirname(__file__), "planning_system.db")

def init_database():
    print(f"Connecting to database at {DB_PATH}...")
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    cursor.execute("PRAGMA foreign_keys = ON;")

    # Drop existing tables
    cursor.executescript("""
    DROP TABLE IF EXISTS wastage_entries;
    DROP TABLE IF EXISTS mrr_entries;
    DROP TABLE IF EXISTS store_transactions;
    DROP TABLE IF EXISTS requisition_items;
    DROP TABLE IF EXISTS material_requisitions;
    DROP TABLE IF EXISTS purchase_requisitions;
    DROP TABLE IF EXISTS production_plans;
    DROP TABLE IF EXISTS sales_forecast;
    DROP TABLE IF EXISTS bom;
    DROP TABLE IF EXISTS materials;
    DROP TABLE IF EXISTS skus;
    DROP TABLE IF EXISTS categories;
    DROP TABLE IF EXISTS erp_settings;
    DROP TABLE IF EXISTS users;
    DROP TABLE IF EXISTS system_settings;

    CREATE TABLE categories (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE NOT NULL, -- e.g. '05', '02', '01'
        name TEXT NOT NULL,         -- e.g. '05. Dry Foods'
        work_center TEXT,          -- Poloxy ERP Work Center name
        branch_id TEXT,            -- e.g. 'B0005', 'B0002', 'B0001'
        description TEXT,
        is_active INTEGER DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        role TEXT NOT NULL, -- Super Admin, Sales, Production, Store, Supply Chain, Management
        department TEXT NOT NULL,
        password TEXT NOT NULL,
        allowed_category TEXT DEFAULT 'ALL', -- 'ALL', '05', '02', '01'
        is_active INTEGER DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE skus (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        sku_code TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        category TEXT,
        sub_category TEXT,
        ctn_size INTEGER DEFAULT 12, -- Number of packets/pieces per carton
        pack_size TEXT,
        uom_default TEXT DEFAULT 'Packet', -- 'Packet', 'Pcs', 'Kg'
        unit TEXT DEFAULT 'Cartoon',
        status TEXT DEFAULT 'Active',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (category_code) REFERENCES categories(code) ON DELETE RESTRICT
    );

    CREATE TABLE materials (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        item_code TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        type TEXT NOT NULL DEFAULT 'RM', -- 'RM' (Raw Material) or 'PM' (Packaging Material)
        category TEXT DEFAULT 'General',
        unit TEXT NOT NULL DEFAULT 'Kg',
        unit_price REAL DEFAULT 0.0,
        current_stock REAL DEFAULT 0.0,
        safety_stock REAL DEFAULT 0.0,
        in_transit REAL DEFAULT 0.0,
        supplier_name TEXT,
        lead_time_days INTEGER DEFAULT 7,
        priority TEXT DEFAULT 'Medium', -- Low, Medium, High
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (category_code) REFERENCES categories(code) ON DELETE RESTRICT
    );

    CREATE TABLE bom (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sku_id INTEGER NOT NULL,
        material_id INTEGER NOT NULL,
        rate_per_ctn REAL NOT NULL DEFAULT 0.0,
        rate_per_packet REAL DEFAULT 0.0,
        unit TEXT DEFAULT 'Kg',
        notes TEXT,
        FOREIGN KEY (sku_id) REFERENCES skus(id) ON DELETE CASCADE,
        FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE CASCADE,
        UNIQUE (sku_id, material_id)
    );

    CREATE TABLE sales_forecast (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sku_id INTEGER NOT NULL,
        month_year TEXT NOT NULL, -- Format: YYYY-MM e.g. 2026-07
        forecast_ctn REAL NOT NULL DEFAULT 0.0,
        forecast_pcs REAL DEFAULT 0.0,
        created_by TEXT DEFAULT 'Sales Team',
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (sku_id) REFERENCES skus(id) ON DELETE CASCADE,
        UNIQUE (sku_id, month_year)
    );

    CREATE TABLE production_plans (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        plan_no TEXT UNIQUE NOT NULL,
        month_year TEXT NOT NULL,
        plan_date TEXT NOT NULL,
        sku_id INTEGER NOT NULL,
        planned_ctn REAL NOT NULL,
        planned_pcs REAL DEFAULT 0.0,
        actual_produced_ctn REAL DEFAULT 0.0,
        actual_produced_pcs REAL DEFAULT 0.0,
        suggested_ctn REAL DEFAULT 0.0,
        status TEXT DEFAULT 'Planned', -- Planned, Requisition Raised, In Production, Completed, Cancelled
        created_by TEXT DEFAULT 'Production Planner',
        notes TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (sku_id) REFERENCES skus(id) ON DELETE RESTRICT
    );

    CREATE TABLE material_requisitions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        req_no TEXT UNIQUE NOT NULL,
        plan_id INTEGER,
        month_year TEXT NOT NULL,
        req_date TEXT NOT NULL,
        status TEXT DEFAULT 'Pending', -- Pending, Approved, Issued, Partially Issued, Rejected
        requested_by TEXT DEFAULT 'Production Team',
        issued_by TEXT,
        issued_date TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (plan_id) REFERENCES production_plans(id) ON DELETE SET NULL
    );

    CREATE TABLE requisition_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        requisition_id INTEGER NOT NULL,
        material_id INTEGER NOT NULL,
        required_qty REAL NOT NULL,
        issued_qty REAL DEFAULT 0.0,
        unit TEXT,
        status TEXT DEFAULT 'Pending',
        FOREIGN KEY (requisition_id) REFERENCES material_requisitions(id) ON DELETE CASCADE,
        FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE RESTRICT
    );

    CREATE TABLE store_transactions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        material_id INTEGER NOT NULL,
        txn_type TEXT NOT NULL, -- INITIAL, ISSUE_PRODUCTION, RECEIPT_GRN, MRR_ENTRY, WASTAGE, ADJUSTMENT
        qty REAL NOT NULL,
        balance_after REAL NOT NULL,
        reference_no TEXT,
        notes TEXT,
        created_by TEXT DEFAULT 'Store Manager',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE CASCADE
    );

    CREATE TABLE mrr_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        mrr_no TEXT UNIQUE NOT NULL,
        mrr_date TEXT NOT NULL,
        supplier_name TEXT NOT NULL,
        chalan_no TEXT,
        material_id INTEGER NOT NULL,
        received_qty REAL NOT NULL,
        unit_price REAL DEFAULT 0.0,
        total_amount REAL DEFAULT 0.0,
        vehicle_no TEXT,
        remarks TEXT,
        created_by TEXT DEFAULT 'Store In-Charge',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE RESTRICT
    );

    CREATE TABLE wastage_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        date TEXT NOT NULL,
        item_type TEXT NOT NULL, -- 'RM', 'PM', 'FG'
        item_id INTEGER NOT NULL,
        batch_no TEXT,
        qty REAL NOT NULL,
        unit TEXT,
        reason TEXT, -- Process Loss, Machine Jam, Quality Rejection, Expired, Handling Damage
        estimated_cost REAL DEFAULT 0.0,
        remarks TEXT,
        created_by TEXT DEFAULT 'Production/Store',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE purchase_requisitions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_code TEXT NOT NULL DEFAULT '05',
        po_no TEXT UNIQUE NOT NULL,
        month_year TEXT NOT NULL,
        material_id INTEGER NOT NULL,
        req_qty REAL NOT NULL,
        unit_price REAL DEFAULT 0.0,
        total_value REAL DEFAULT 0.0,
        supplier_name TEXT,
        lead_time_days INTEGER DEFAULT 7,
        priority TEXT DEFAULT 'Medium',
        status TEXT DEFAULT 'Pending', -- Pending, Approved, PO Issued, In-Transit, Received, Cancelled
        ordered_date TEXT,
        expected_delivery TEXT,
        received_qty REAL DEFAULT 0.0,
        created_by TEXT DEFAULT 'Store Department',
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE RESTRICT
    );

    CREATE TABLE erp_settings (
        key TEXT PRIMARY KEY,
        value TEXT,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE system_settings (
        key TEXT PRIMARY KEY,
        value TEXT
    );
    """)

    # 1. Seed Product Categories
    categories = [
        ('05', '05. Dry Foods', 'Dry Food Packing(Bakery Puffed Sauce Chanachur & Lachsa)', 'B0005', 'Bakery, Cakes, Biscuits, Chanachur & Puffed Rice'),
        ('02', '02. Frozen Foods', 'Frozen Finished Goods', 'B0002', 'Nuggets, Sausages, Paratha, Samosa & Spring Rolls'),
        ('01', '01. Process Chicken', 'Process Chicken Plant', 'B0001', 'Fresh & Chilled Broiler Chicken, Breast, Drumsticks')
    ]
    cursor.executemany("""
        INSERT INTO categories (code, name, work_center, branch_id, description)
        VALUES (?, ?, ?, ?, ?)
    """, categories)

    # 2. Seed Default ERP Credentials
    erp_defaults = [
        ('poloxy_url', 'http://erp.paragon.com.bd'),
        ('poloxy_user', 'a0162e0019'),
        ('poloxy_pass', 'pal.123456'),
        ('default_uom_display', 'Packet'), # 'Packet' or 'Carton'
        ('dry_food_work_center', 'Dry Food Packing(Bakery Puffed Sauce Chanachur & Lachsa)'),
        ('frozen_food_work_center', 'Frozen Finished Goods'),
        ('chicken_work_center', 'Process Chicken Plant')
    ]
    cursor.executemany("INSERT INTO erp_settings (key, value) VALUES (?, ?)", erp_defaults)

    # 3. Seed Users with Category Access Rights
    users_data = [
        ('super_admin', 'Super Admin (System Manager)', 'Super Admin', 'Management & IT', 'admin123', 'ALL'),
        ('admin', 'Paragon Plant Admin (Viewer & Reports)', 'Admin', 'Plant Administration', 'admin123', 'ALL'),
        ('dry_sales', 'Md. Rafiqul Islam (Dry Food Sales)', 'Sales', 'Sales & Marketing', 'sales123', '05'),
        ('frozen_sales', 'Tanvir Hasan (Frozen Food Sales)', 'Sales', 'Sales & Marketing', 'sales123', '02'),
        ('chicken_sales', 'Enamul Haque (Chicken Sales)', 'Sales', 'Sales & Marketing', 'sales123', '01'),
        ('dry_prod', 'Engr. Kamal Hossain (Dry Food Plant)', 'Production', 'Production & Planning', 'prod123', '05'),
        ('frozen_prod', 'Engr. Sohel Rana (Frozen Plant)', 'Production', 'Production & Planning', 'prod123', '02'),
        ('dry_store', 'Md. Tariqul Anam (Dry Food Store)', 'Store', 'Store & Inventory', 'store123', '05'),
        ('frozen_store', 'Zahirul Islam (Cold Storage Store)', 'Store', 'Cold Chain Store', 'store123', '02'),
        ('supplychain', 'Tanvir Ahmed (Procurement Lead)', 'Supply Chain', 'Procurement & SCM', 'scm123', 'ALL'),
        ('management', 'Director / GM Sir (Executive)', 'Management', 'Executive Board', 'mgmt123', 'ALL')
    ]
    cursor.executemany("""
        INSERT INTO users (username, name, role, department, password, allowed_category)
        VALUES (?, ?, ?, ?, ?, ?)
    """, users_data)

    print("Loading Excel workbook data for Dry Foods (05)...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    po_sheet = wb['📋 PO Summary']
    po_info = {}
    for r in range(4, po_sheet.max_row + 1):
        name = po_sheet.cell(r, 2).value
        if name and not str(name).startswith('MONTH GUIDE'):
            name_clean = str(name).strip()
            up = po_sheet.cell(r, 5).value or 0
            lt = po_sheet.cell(r, 7).value or 7
            prio = po_sheet.cell(r, 9).value or 'Medium'
            po_info[name_clean] = {
                'unit_price': float(up),
                'lead_time': int(lt) if isinstance(lt, (int, float)) else 7,
                'priority': str(prio).strip()
            }

    store_sheet = wb['🏪 Store & Requisition']
    material_map = {}
    mat_code_idx = 1

    for r in range(4, store_sheet.max_row + 1):
        name = store_sheet.cell(r, 1).value
        if not name or str(name).startswith('MONTH GUIDE'):
            continue
        supplier = str(store_sheet.cell(r, 10).value or 'Standard Supplier').strip()
        name_clean = str(name).strip()
        if name_clean in material_map:
            name_clean = f"{name_clean} ({supplier})"

        unit = str(store_sheet.cell(r, 2).value or 'Kg').strip()
        op_stock = float(store_sheet.cell(r, 4).value or 0)
        safety_stock = float(store_sheet.cell(r, 5).value or 0)
        in_transit = float(store_sheet.cell(r, 6).value or 0)
        lead_time = store_sheet.cell(r, 11).value or 7
        lead_time = int(lead_time) if isinstance(lead_time, (int, float)) else 7

        p_info = po_info.get(name_clean, {})
        unit_price = p_info.get('unit_price', 50.0)
        priority = p_info.get('priority', 'Medium')
        item_code = f"RM-{mat_code_idx:04d}"
        mat_code_idx += 1

        cursor.execute("""
            INSERT INTO materials (category_code, item_code, name, type, category, unit, unit_price, current_stock, safety_stock, in_transit, supplier_name, lead_time_days, priority)
            VALUES ('05', ?, ?, 'RM', 'Raw Material', ?, ?, ?, ?, ?, ?, ?, ?)
        """, (item_code, name_clean, unit, unit_price, op_stock, safety_stock, in_transit, supplier, lead_time, priority))
        material_map[name_clean] = cursor.lastrowid

        cursor.execute("""
            INSERT INTO store_transactions (material_id, txn_type, qty, balance_after, reference_no, notes, created_by)
            VALUES (?, 'INITIAL', ?, ?, 'INIT-OPENING', 'Opening balance imported from Excel plan', 'System Init')
        """, (material_map[name_clean], op_stock, op_stock))

    # Add Dry Food PM items
    pm_list = [
        ("PM-0001", "Master Outer Carton (5-Ply Corrugated)", "Pcs", 38.5, 8500, 2500, 1000, "Bengal Cartons Ltd", 5, "High"),
        ("PM-0002", "Printed Inner Box (Duplex Board 350 GSM)", "Pcs", 8.2, 24000, 6000, 0, "National Packaging", 7, "High"),
        ("PM-0003", "Metalized Barrier Foil (Roll)", "Kg", 320.0, 450, 120, 50, "Standard Poly & Foil", 10, "High"),
        ("PM-0004", "BOPP Clear Film 25 Micron (Roll)", "Kg", 260.0, 520, 150, 0, "Standard Poly & Foil", 7, "Medium"),
        ("PM-0005", "Shrink Film Poly Wrap", "Kg", 185.0, 310, 80, 0, "Dhaka Plastic Works", 4, "Medium"),
        ("PM-0006", "Sealing BOPP Brown Tape (72mm)", "Roll", 125.0, 650, 150, 0, "Star Adhesive & Tape", 3, "Medium"),
        ("PM-0007", "Corrugated Partition / Layer Separator", "Pcs", 3.5, 12000, 3000, 0, "Bengal Cartons Ltd", 5, "Low"),
        ("PM-0008", "Printed Batch / MRP / Date Sticker", "Pcs", 0.45, 85000, 20000, 0, "Quick Print Solutions", 3, "High"),
        ("PM-0009", "Silica Gel Desiccant Pouch (1g)", "Pcs", 1.2, 15000, 4000, 0, "Chem-BD Imports", 15, "Low"),
        ("PM-0010", "Plastic Tray Inner Insert (PET)", "Pcs", 4.8, 18000, 5000, 2000, "Modern Thermoforming", 10, "Medium")
    ]
    for pm_code, pm_name, pm_unit, pm_price, pm_stock, pm_safety, pm_transit, pm_supp, pm_lt, pm_prio in pm_list:
        cursor.execute("""
            INSERT INTO materials (category_code, item_code, name, type, category, unit, unit_price, current_stock, safety_stock, in_transit, supplier_name, lead_time_days, priority)
            VALUES ('05', ?, ?, 'PM', 'Packaging Material', ?, ?, ?, ?, ?, ?, ?, ?)
        """, (pm_code, pm_name, pm_unit, pm_price, pm_stock, pm_safety, pm_transit, pm_supp, pm_lt, pm_prio))
        material_map[pm_name] = cursor.lastrowid
        cursor.execute("""
            INSERT INTO store_transactions (material_id, txn_type, qty, balance_after, reference_no, notes, created_by)
            VALUES (?, 'INITIAL', ?, ?, 'INIT-OPENING', 'Initial PM stock registered', 'System Init')
        """, (material_map[pm_name], pm_stock, pm_stock))

    # Add Frozen Foods (02) Materials & SKUs
    frozen_mats = [
        ("02", "RM-FZ-001", "Deboned Chicken Meat (Frozen)", "RM", "Kg", 240.0, 4500.0, 1500.0, 500.0, "Paragon Breeder & Hatchery", 3, "High"),
        ("02", "RM-FZ-002", "Golden Breadcrumbs Coating", "RM", "Kg", 110.0, 1800.0, 500.0, 0.0, "Bake BD Supplies", 5, "High"),
        ("02", "RM-FZ-003", "Batter Mix Powder (Adhesion)", "RM", "Kg", 95.0, 1200.0, 400.0, 0.0, "Agro Ingredient Ltd", 7, "Medium"),
        ("02", "RM-FZ-004", "Seasoning Spice Blend (Nugget/Sausage)", "RM", "Kg", 320.0, 450.0, 150.0, 0.0, "Taste Master BD", 10, "High"),
        ("02", "RM-FZ-005", "Wheat Flour (Special Paratha Grade)", "RM", "Kg", 58.0, 6000.0, 2000.0, 0.0, "City Millers Ltd", 4, "High"),
        ("02", "RM-FZ-006", "Refined Vegetable Ghee / Shortening", "RM", "Kg", 145.0, 1800.0, 600.0, 0.0, "Delta Oils & Fats", 3, "Medium"),
        ("02", "PM-FZ-001", "Printed 3-Layer Vacuum Pouch 250g", "PM", "Pcs", 4.2, 35000.0, 10000.0, 5000.0, "Standard Poly & Foil", 7, "High"),
        ("02", "PM-FZ-002", "Frozen Master Carton (Cold-Resistant 5-Ply)", "PM", "Pcs", 42.0, 4500.0, 1500.0, 0.0, "Bengal Cartons Ltd", 5, "High"),
        ("02", "PM-FZ-003", "Interleaving Plastic Sheet (Paratha Separator)", "PM", "Pcs", 0.65, 80000.0, 20000.0, 0.0, "Dhaka Plastic Works", 4, "Medium")
    ]
    for f_cat, f_code, f_name, f_type, f_unit, f_price, f_stk, f_safe, f_tr, f_supp, f_lt, f_prio in frozen_mats:
        cursor.execute("""
            INSERT INTO materials (category_code, item_code, name, type, category, unit, unit_price, current_stock, safety_stock, in_transit, supplier_name, lead_time_days, priority)
            VALUES (?, ?, ?, ?, 'Frozen Materials', ?, ?, ?, ?, ?, ?, ?, ?)
        """, (f_cat, f_code, f_name, f_type, f_unit, f_price, f_stk, f_safe, f_tr, f_supp, f_lt, f_prio))
        material_map[f_name] = cursor.lastrowid
        cursor.execute("""
            INSERT INTO store_transactions (material_id, txn_type, qty, balance_after, reference_no, notes, created_by)
            VALUES (?, 'INITIAL', ?, ?, 'INIT-OPENING', 'Initial Frozen material stock', 'System Init')
        """, (material_map[f_name], f_stk, f_stk))

    # Add Process Chicken (01) Materials
    chicken_mats = [
        ("01", "RM-CK-001", "Live Broiler Bird (Average 1.6 Kg)", "RM", "Kg", 165.0, 12000.0, 3000.0, 0.0, "Paragon Broiler Farm Gazipur", 1, "High"),
        ("01", "PM-CK-001", "Printed Poly Bag with Ventilation (Whole)", "PM", "Pcs", 2.8, 25000.0, 8000.0, 0.0, "Standard Poly & Foil", 5, "High"),
        ("01", "PM-CK-002", "Plastic Crate (Reusable 20 Kg)", "PM", "Pcs", 350.0, 800.0, 200.0, 0.0, "Dhaka Plastic Works", 10, "Low"),
        ("01", "PM-CK-003", "Styrofoam Tray with Cling Film (Cut-up)", "PM", "Pcs", 3.5, 15000.0, 5000.0, 0.0, "Modern Thermoforming", 4, "Medium")
    ]
    for c_cat, c_code, c_name, c_type, c_unit, c_price, c_stk, c_safe, c_tr, c_supp, c_lt, c_prio in chicken_mats:
        cursor.execute("""
            INSERT INTO materials (category_code, item_code, name, type, category, unit, unit_price, current_stock, safety_stock, in_transit, supplier_name, lead_time_days, priority)
            VALUES (?, ?, ?, ?, 'Poultry Materials', ?, ?, ?, ?, ?, ?, ?, ?)
        """, (c_cat, c_code, c_name, c_type, c_unit, c_price, c_stk, c_safe, c_tr, c_supp, c_lt, c_prio))
        material_map[c_name] = cursor.lastrowid

    # Load SKUs from Excel (Dry Foods)
    sku_sheet = wb['📝 SKU Master']
    sku_map = {}
    for r in range(4, sku_sheet.max_row + 1):
        sid = sku_sheet.cell(r, 2).value
        sname = sku_sheet.cell(r, 3).value
        if not sid or not sname:
            continue
        sid_clean = str(sid).strip()
        sname_clean = str(sname).strip()
        cat = str(sku_sheet.cell(r, 4).value or 'Cookies & Cake').strip()
        subcat = str(sku_sheet.cell(r, 5).value or 'Standard').strip()
        ctn_size = int(sku_sheet.cell(r, 6).value or 12)
        pack_size = str(sku_sheet.cell(r, 7).value or '').strip()

        cursor.execute("""
            INSERT INTO skus (category_code, sku_code, name, category, sub_category, ctn_size, pack_size, uom_default, unit, status)
            VALUES ('05', ?, ?, ?, ?, ?, ?, 'Packet', 'Cartoon', 'Active')
        """, (sid_clean, sname_clean, cat, subcat, ctn_size, pack_size))
        sku_id = cursor.lastrowid
        sku_map[sid_clean] = sku_id
        sku_map[sname_clean] = sku_id

    # Add Frozen Foods SKUs
    frozen_skus = [
        ("02", "SKU-FZ-001", "Paragon Chicken Nuggets 250gm", "Frozen Snack", "Value Added", 24, "250gm", "Packet"),
        ("02", "SKU-FZ-002", "Paragon Chicken Nuggets 500gm", "Frozen Snack", "Value Added", 12, "500gm", "Packet"),
        ("02", "SKU-FZ-003", "Paragon Chicken Sausage 250gm", "Frozen Snack", "Breakfast", 24, "250gm", "Packet"),
        ("02", "SKU-FZ-004", "Paragon Plain Paratha (10 Pcs Pack)", "Frozen Bread", "Traditional", 16, "800gm", "Packet"),
        ("02", "SKU-FZ-005", "Paragon Chicken Samosa (10 Pcs Pack)", "Frozen Snack", "Snack", 20, "300gm", "Packet"),
        ("02", "SKU-FZ-006", "Paragon Chicken Singara (10 Pcs Pack)", "Frozen Snack", "Snack", 20, "350gm", "Packet")
    ]
    for fz_cat, fz_code, fz_name, fz_catname, fz_sub, fz_ctn, fz_pack, fz_uom in frozen_skus:
        cursor.execute("""
            INSERT INTO skus (category_code, sku_code, name, category, sub_category, ctn_size, pack_size, uom_default, unit, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'Cartoon', 'Active')
        """, (fz_cat, fz_code, fz_name, fz_catname, fz_sub, fz_ctn, fz_pack, fz_uom))
        s_id = cursor.lastrowid
        sku_map[fz_code] = s_id
        sku_map[fz_name] = s_id

        # BOM for Frozen SKUs
        if "Nuggets" in fz_name:
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 3.6, 'Kg')", (s_id, material_map["Deboned Chicken Meat (Frozen)"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 1.2, 'Kg')", (s_id, material_map["Golden Breadcrumbs Coating"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 0.8, 'Kg')", (s_id, material_map["Batter Mix Powder (Adhesion)"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 24.0, 'Pcs')", (s_id, material_map["Printed 3-Layer Vacuum Pouch 250g"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 1.0, 'Pcs')", (s_id, material_map["Frozen Master Carton (Cold-Resistant 5-Ply)"]))
        elif "Paratha" in fz_name:
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 9.5, 'Kg')", (s_id, material_map["Wheat Flour (Special Paratha Grade)"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 2.2, 'Kg')", (s_id, material_map["Refined Vegetable Ghee / Shortening"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 160.0, 'Pcs')", (s_id, material_map["Interleaving Plastic Sheet (Paratha Separator)"]))
            cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 1.0, 'Pcs')", (s_id, material_map["Frozen Master Carton (Cold-Resistant 5-Ply)"]))

    # Add Process Chicken SKUs
    chicken_skus = [
        ("01", "SKU-CK-001", "Paragon Whole Broiler (Skin-On Chilled)", "Dressed Chicken", "Fresh", 10, "1.2Kg", "Pcs"),
        ("01", "SKU-CK-002", "Paragon Chicken Breast Boneless", "Cut-Up Parts", "Prime", 12, "1Kg", "Packet"),
        ("01", "SKU-CK-003", "Paragon Chicken Drumstick (Fresh)", "Cut-Up Parts", "Prime", 12, "1Kg", "Packet")
    ]
    for ck_cat, ck_code, ck_name, ck_catname, ck_sub, ck_ctn, ck_pack, ck_uom in chicken_skus:
        cursor.execute("""
            INSERT INTO skus (category_code, sku_code, name, category, sub_category, ctn_size, pack_size, uom_default, unit, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'Crate', 'Active')
        """, (ck_cat, ck_code, ck_name, ck_catname, ck_sub, ck_ctn, ck_pack, ck_uom))
        s_id = cursor.lastrowid
        sku_map[ck_code] = s_id
        sku_map[ck_name] = s_id
        cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 16.5, 'Kg')", (s_id, material_map["Live Broiler Bird (Average 1.6 Kg)"]))
        cursor.execute("INSERT INTO bom (sku_id, material_id, rate_per_ctn, unit) VALUES (?, ?, 10.0, 'Pcs')", (s_id, material_map["Printed Poly Bag with Ventilation (Whole)"]))

    # Load BOM from Excel for Dry Foods
    bom_sheet = wb['⚗️ RM Consumption']
    col_mat_map = {}
    for c in range(8, bom_sheet.max_column + 1):
        m_name = bom_sheet.cell(2, c).value
        if m_name:
            m_clean = str(m_name).strip()
            if m_clean in material_map:
                col_mat_map[c] = material_map[m_clean]

    for r in range(4, bom_sheet.max_row + 1):
        s_name = bom_sheet.cell(r, 4).value
        if not s_name:
            continue
        s_clean = str(s_name).strip()
        sku_id = sku_map.get(s_clean)
        if not sku_id:
            code_val = bom_sheet.cell(r, 2).value
            if code_val:
                sku_id = sku_map.get(f"SKU-{int(code_val):03d}")
        if not sku_id:
            continue

        for c, m_id in col_mat_map.items():
            rate = bom_sheet.cell(r, c).value
            if rate and isinstance(rate, (int, float)) and rate > 0:
                cursor.execute("""
                    INSERT OR REPLACE INTO bom (sku_id, material_id, rate_per_ctn, unit)
                    VALUES (?, ?, ?, 'Kg')
                """, (sku_id, m_id, float(rate)))

        # Default Dry Food PM in BOM
        carton_mat_id = material_map.get("Master Outer Carton (5-Ply Corrugated)")
        tape_mat_id = material_map.get("Sealing BOPP Brown Tape (72mm)")
        sticker_mat_id = material_map.get("Printed Batch / MRP / Date Sticker")
        foil_mat_id = material_map.get("Metalized Barrier Foil (Roll)")

        if carton_mat_id:
            cursor.execute("INSERT OR IGNORE INTO bom (sku_id, material_id, rate_per_ctn, unit, notes) VALUES (?, ?, 1.0, 'Pcs', '1 Outer Carton per CTN')", (sku_id, carton_mat_id))
        if tape_mat_id:
            cursor.execute("INSERT OR IGNORE INTO bom (sku_id, material_id, rate_per_ctn, unit, notes) VALUES (?, ?, 0.05, 'Roll', '0.05 Roll Tape per CTN')", (sku_id, tape_mat_id))
        if sticker_mat_id:
            cursor.execute("INSERT OR IGNORE INTO bom (sku_id, material_id, rate_per_ctn, unit, notes) VALUES (?, ?, 2.0, 'Pcs', '2 Labels per CTN')", (sku_id, sticker_mat_id))
        if foil_mat_id:
            cursor.execute("INSERT OR IGNORE INTO bom (sku_id, material_id, rate_per_ctn, unit, notes) VALUES (?, ?, 0.12, 'Kg', 'Foil packing per CTN')", (sku_id, foil_mat_id))

    # Load Sales Forecast
    fc_sheet = wb['📅 Sales Forecast']
    month_cols = []
    for c in range(5, 17):
        val = fc_sheet.cell(3, c).value
        if isinstance(val, datetime):
            m_str = val.strftime("%Y-%m")
            month_cols.append((c, m_str))
        elif val:
            month_cols.append((c, str(val)[:7]))

    for r in range(4, fc_sheet.max_row + 1):
        sid = fc_sheet.cell(r, 2).value
        sname = fc_sheet.cell(r, 3).value
        sku_id = None
        if sid:
            sku_id = sku_map.get(str(sid).strip())
        if not sku_id and sname:
            sku_id = sku_map.get(str(sname).strip())
        if not sku_id:
            continue

        for c, m_str in month_cols:
            val = fc_sheet.cell(r, c).value
            if val and isinstance(val, (int, float)) and val > 0:
                cursor.execute("""
                    INSERT OR REPLACE INTO sales_forecast (sku_id, month_year, forecast_ctn)
                    VALUES (?, ?, ?)
                """, (sku_id, m_str, float(val)))

    # Seed Frozen & Chicken Forecasts for July 2026
    for fz_code in ["SKU-FZ-001", "SKU-FZ-002", "SKU-FZ-003", "SKU-FZ-004"]:
        f_id = sku_map.get(fz_code)
        if f_id:
            cursor.execute("INSERT OR REPLACE INTO sales_forecast (sku_id, month_year, forecast_ctn) VALUES (?, '2026-07', 1500)", (f_id,))
    for ck_code in ["SKU-CK-001", "SKU-CK-002"]:
        c_id = sku_map.get(ck_code)
        if c_id:
            cursor.execute("INSERT OR REPLACE INTO sales_forecast (sku_id, month_year, forecast_ctn) VALUES (?, '2026-07', 2200)", (c_id,))

    # Seed Demo Production Plans
    cursor.execute("""
        SELECT sf.sku_id, s.category_code, s.sku_code, s.name, sf.forecast_ctn, s.ctn_size
        FROM sales_forecast sf 
        JOIN skus s ON sf.sku_id = s.id 
        WHERE sf.month_year = '2026-07' AND sf.forecast_ctn > 0
        LIMIT 10
    """)
    active_plans = cursor.fetchall()
    plan_idx = 1
    for sku_id, cat_code, scode, sname, fc_ctn, ctn_size in active_plans:
        plan_no = f"PP-{cat_code}-202607-{plan_idx:03d}"
        planned_ctn = fc_ctn
        planned_pcs = planned_ctn * ctn_size
        actual_produced = round(planned_ctn * (0.85 if plan_idx % 2 == 0 else 0.70), 0)
        actual_pcs = actual_produced * ctn_size
        status = 'In Production' if plan_idx <= 4 else 'Completed'

        cursor.execute("""
            INSERT INTO production_plans (category_code, plan_no, month_year, plan_date, sku_id, planned_ctn, planned_pcs, actual_produced_ctn, actual_produced_pcs, suggested_ctn, status, notes)
            VALUES (?, ?, '2026-07', '2026-07-02', ?, ?, ?, ?, ?, ?, ?, 'Monthly production run against forecast')
        """, (cat_code, plan_no, sku_id, planned_ctn, planned_pcs, actual_produced, actual_pcs, planned_ctn, status))
        plan_id = cursor.lastrowid

        req_no = f"MR-{cat_code}-202607-{plan_idx:03d}"
        req_status = 'Issued' if status == 'Completed' else 'Partially Issued'
        cursor.execute("""
            INSERT INTO material_requisitions (category_code, req_no, plan_id, month_year, req_date, status, requested_by, issued_by, issued_date, remarks)
            VALUES (?, ?, ?, '2026-07', '2026-07-03', ?, 'Engr. Kamal Hossain', 'Md. Tariqul Anam', '2026-07-03', 'Materials issued for production batch')
        """, (cat_code, req_no, plan_id, req_status))
        req_id = cursor.lastrowid

        cursor.execute("SELECT material_id, rate_per_ctn, unit FROM bom WHERE sku_id = ?", (sku_id,))
        bom_rows = cursor.fetchall()
        for m_id, rate, m_unit in bom_rows:
            req_qty = round(planned_ctn * rate, 2)
            issued_qty = req_qty if req_status == 'Issued' else round(req_qty * 0.8, 2)
            cursor.execute("""
                INSERT INTO requisition_items (requisition_id, material_id, required_qty, issued_qty, unit, status)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (req_id, m_id, req_qty, issued_qty, m_unit, req_status))

            cursor.execute("SELECT current_stock, name FROM materials WHERE id = ?", (m_id,))
            cur_mat = cursor.fetchone()
            if cur_mat:
                curr_stk, mat_name = cur_mat
                new_stk = max(0.0, curr_stk - issued_qty)
                cursor.execute("UPDATE materials SET current_stock = ? WHERE id = ?", (new_stk, m_id))
                cursor.execute("""
                    INSERT INTO store_transactions (material_id, txn_type, qty, balance_after, reference_no, notes, created_by)
                    VALUES (?, 'ISSUE_PRODUCTION', ?, ?, ?, ?, 'Md. Tariqul Anam')
                """, (m_id, -issued_qty, new_stk, req_no, f"Issued to Production for Plan {plan_no} ({sname})"))

        plan_idx += 1

    # Seed Sample MRR (Material Receive Reports) for Store Stock Entry
    sample_mrr = [
        ('05', 'MRR-202607-001', '2026-07-04', 'ABC Trading', 'CH-9921', material_map.get('Flour', 1), 2000.0, 45.0, 90000.0, 'Dhaka Metro Ta-11-2091', 'Raw wheat flour received in good condition'),
        ('05', 'MRR-202607-002', '2026-07-05', 'Bengal Cartons Ltd', 'CH-4482', material_map.get('Master Outer Carton (5-Ply Corrugated)', 2), 5000.0, 38.5, 192500.0, 'Dhaka Metro Cha-53-1180', 'Outer corrugated cartons received and stacked'),
        ('02', 'MRR-202607-003', '2026-07-06', 'Paragon Breeder & Hatchery', 'CH-8819', material_map.get('Deboned Chicken Meat (Frozen)', 3), 3000.0, 240.0, 720000.0, 'Reefer Van DM-7712', 'Frozen chicken meat received at -18C temp')
    ]
    cursor.executemany("""
        INSERT INTO mrr_entries (category_code, mrr_no, mrr_date, supplier_name, chalan_no, material_id, received_qty, unit_price, total_amount, vehicle_no, remarks)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, sample_mrr)

    # Seed Sample Wastage Entries
    sample_wastage = [
        ('05', '2026-07-08', 'RM', material_map.get('Sugar', 1), 'B-202607-1', 15.5, 'Kg', 'Process Loss', 186.0, 'Spilled during dough mixing hopper feed'),
        ('05', '2026-07-09', 'PM', material_map.get('Printed Inner Box (Duplex Board 350 GSM)', 2), 'B-202607-2', 45.0, 'Pcs', 'Machine Jam', 369.0, 'Packing line sensor misalignment jam'),
        ('02', '2026-07-10', 'RM', material_map.get('Golden Breadcrumbs Coating', 3), 'B-FZ-001', 12.0, 'Kg', 'Handling Damage', 1320.0, 'Torn bag in transit inside cold store')
    ]
    cursor.executemany("""
        INSERT INTO wastage_entries (category_code, date, item_type, item_id, batch_no, qty, unit, reason, estimated_cost, remarks)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, sample_wastage)

    # Seed Purchase Requisitions
    cursor.execute("""
        SELECT id, category_code, name, unit, current_stock, safety_stock, unit_price, supplier_name, lead_time_days, priority 
        FROM materials 
        WHERE current_stock <= (safety_stock * 1.5)
        LIMIT 8
    """)
    low_stock_items = cursor.fetchall()
    po_idx = 1
    for m_id, cat_code, m_name, m_unit, c_stk, s_stk, u_price, supp, lt, prio in low_stock_items:
        deficit = max(s_stk * 2, (s_stk - c_stk) + (s_stk * 1.5))
        req_qty = round(deficit, 1)
        tot_val = round(req_qty * u_price, 2)
        po_no = f"PR-{cat_code}-202607-{po_idx:03d}"
        status = 'In-Transit' if po_idx <= 2 else 'Pending'

        cursor.execute("""
            INSERT INTO purchase_requisitions (category_code, po_no, month_year, material_id, req_qty, unit_price, total_value, supplier_name, lead_time_days, priority, status, ordered_date, expected_delivery, created_by, remarks)
            VALUES (?, ?, '2026-07', ?, ?, ?, ?, ?, ?, ?, ?, '2026-07-05', '2026-07-15', 'Store In-Charge', 'Requisition raised due to low stock / safety stock threshold')
        """, (cat_code, po_no, m_id, req_qty, u_price, tot_val, supp, lt, prio, status))

        if status == 'In-Transit':
            cursor.execute("UPDATE materials SET in_transit = in_transit + ? WHERE id = ?", (req_qty, m_id))

        po_idx += 1

    conn.commit()
    conn.close()
    print("Multi-category Database initialization and seed completed successfully!")

if __name__ == '__main__':
    init_database()
